#!/usr/bin/env python3
"""Create a separate editable figure-enhanced DOCX; never overwrite the input."""
import argparse,copy,json,re,hashlib
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Inches,Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT=Path(__file__).resolve().parents[2]

def para_after(anchor,text='',style=None):
 p=OxmlElement('w:p');anchor._p.addnext(p);o=Paragraph(p,anchor._parent)
 if style:o.style=style
 if text:o.add_run(text)
 return o

def field(p,code,display=''):
 r=OxmlElement('w:r');beg=OxmlElement('w:fldChar');beg.set(qn('w:fldCharType'),'begin');beg.set(qn('w:dirty'),'true');r.append(beg);p._p.append(r)
 r=OxmlElement('w:r');t=OxmlElement('w:instrText');t.set(qn('xml:space'),'preserve');t.text=' '+code+' ';r.append(t);p._p.append(r)
 if display:
  r=OxmlElement('w:r');s=OxmlElement('w:fldChar');s.set(qn('w:fldCharType'),'separate');r.append(s);p._p.append(r)
  for j,line in enumerate(display.split('\n')):
   if j:p.add_run().add_break()
   p.add_run(line)
 r=OxmlElement('w:r');e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end');r.append(e);p._p.append(r)

def tc(p,text,kind,level=1):
 # TC fields are hidden in Word. They do not add visible prose to the thesis.
 safe=text.replace('"',"'")
 field(p,f'TC "{safe}" \\f {kind} \\l {level}')
 for r in list(p._p)[-3:]:
  if r.tag==qn('w:r'):
   rp=r.find(qn('w:rPr'))
   if rp is None:rp=OxmlElement('w:rPr');r.insert(0,rp)
   rp.append(OxmlElement('w:vanish'))

def remove(p):
 parent=p._p.getparent()
 if parent is not None:parent.remove(p._p)

def new_table(doc,anchor,no,payload,width):
 p=para_after(anchor,f'Table {no}: {payload["caption"]}','PDFCaption');p.paragraph_format.keep_with_next=True;tc(p,f'{no} {payload["caption"]}','T')
 table=doc.add_table(rows=1,cols=len(payload['headers']));table.autofit=False
 # Generous first column for labels; numeric/result columns share the rest.
 n=len(payload['headers']);first=.42 if n>3 else .34
 ratios={'3.2':[.17,.12,.19,.52],'4.1':[.30,.2333,.2333,.2334],'4.2':[.22,.26,.26,.26],'5.3':[.18,.12,.23,.23,.24],'5.4':[.24,.152,.152,.152,.152,.152],'A.4':[.12,.44,.44],'A.8':[.18,.18,.09,.09,.09,.29,.08],'A.9':[.12,.22,.22,.22,.22]}.get(no)
 if ratios:colwidths=[width*x for x in ratios]
 else:colwidths=[width*first]+[width*(1-first)/(n-1)]*(n-1)
 for col,w in zip(table.columns,colwidths):col.width=Inches(w)
 for cell,txt in zip(table.rows[0].cells,payload['headers']):cell.text=str(txt)
 for row in payload['rows']:
  assert len(row)==n,(no,row,n)
  for cell,txt in zip(table.add_row().cells,row):cell.text=str(txt)
 for ri,row in enumerate(table.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if ri==0:
   h=OxmlElement('w:tblHeader');h.set(qn('w:val'),'true');pr.append(h)
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(colwidths[ci])
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAF0F3');cell._tc.get_or_add_tcPr().append(sh)
   for pp in cell.paragraphs:
    pp.paragraph_format.space_before=Pt(2);pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.line_spacing=1.05
    pp.paragraph_format.keep_with_next=False
    for run in pp.runs:
     run.font.name='Times New Roman';run.font.size=Pt(8.7 if n>4 else 9.2);run.bold=ri==0
 borders=OxmlElement('w:tblBorders')
 for side in ('top','bottom','insideH'):
  b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'AAB5BC');borders.append(b)
 table._tbl.tblPr.append(borders)
 p._p.addnext(table._tbl)
 # Return the caption; caller only needs stable order after original anchors.
 return p,table

FIGURES=[
 ('3.3',282,'methodology/fig_3_3_task_storyboard.png','Figure 3.3 provides representative recorded observations for the four coarse contexts. The images show the visual setting and task variation underlying the benchmark. They are examples of dataset execution rather than a qualitative evaluation of the learnt controllers.','Representative head-camera observations from retained validation source 0. Panels show waiting, acquiring the trash can, picking up and depositing a soda can, following the virtual human, and interacting with the radio. Panel identifiers give the round and raw global frame (30 Hz). Human-facing phases use scripted control; manipulation uses recorded BEHAVIOR replay. These are dataset examples, not successful fine-tuned policy rollouts.'),
 ('3.4',294,'methodology/fig_3_4_controller_data_flow.png','Figure 3.4 locates the learnt controller within the implemented HRC system. The policy replaces replay only inside ExecuteBehaviorTask, while the surrounding interaction and transition logic remains scripted. Source identity is retained from recording through export to the diagnostic benchmark.','Implemented controller scope and data flow. (a) The recorded finite-state schedule; non-final rounds continue through WaitNextCommand and Relocate, while the final round exits from FaceHuman to AllTasksDone. (b) ExecuteBehaviorTask uses upstream state/action replay for dataset construction and the predicted-action client for policy rollout. (c) Source-level quality control and splitting precede diagnostic use of the exported data. High-level phases are supplied by the scripted scheduler.'),
 ('3.5',311,'dataset/fig_3_5_dataset_construction.png','Figure 3.5 summarises dataset construction and the supervision available to the learner. Quality control and splitting operate at source level. Chunk validity and weighted sampling operate on derived anchors, so the natural decision-frame population differs from the configured training-sampler target.','Dataset construction and effective supervision. (a) Quality-control outcomes for 65 collected sources and the frozen split of the 45 retained sources. (b) Complete recorded source durations at 30 Hz, including scripted transition phases. (c) Decision frames account for 3.955% of valid B/C training anchors; the 20% value is the configured sampling target, rather than an observed frequency of realised training draws.'),
 ('A.1',802,'experiments/fig_A_1_training_diagnostics.png','Figure A.1 records the optimisation history of the archived runs. It documents training progress and the fixed 20,000-update endpoint; the paper-facing comparison is instead defined by the frozen physical-action audit.','Training diagnostics for the three 20,000-update runs (one training seed, 42, per configuration). The training objective was logged every 100 updates and the monitored validation objective every 1,000 updates. All 32 model output dimensions, including padding, contribute to the normalised flow objective. Different supervision and normalisation pipelines prevent interpreting these magnitudes as directly comparable physical-action errors or task-success measures.')]

def build(inp,out):
 assert inp.resolve()!=out.resolve(),'Input must be preserved.'
 before=hashlib.sha256(inp.read_bytes()).hexdigest();doc=Document(inp);orig=list(doc.paragraphs);assert len(orig)==836
 for section in doc.sections:section.page_width=Cm(21);section.page_height=Cm(29.7)
 width=min((s.page_width-s.left_margin-s.right_margin)/914400 for s in doc.sections)
 for shape in doc.inline_shapes:
  if shape.width>Inches(width):
   ratio=float(Inches(width))/shape.width;shape.height=int(shape.height*ratio);shape.width=Inches(width)
 payload=json.loads((ROOT/'tables/native_tables.json').read_text());assert len(payload)==18
 # Original image/table placeholder paragraphs; retain references before mutations.
 positions={'3.1':294,'3.2':304,'4.1':472,'4.2':518,'5.1':550,'5.2':560,'5.3':570,'5.4':580,'5.5':590,'A.1':789,'A.2':794,'A.3':795,'A.4':804,'A.5':810,'A.6':815,'A.7':820,'A.8':823,'A.9':825}
 table_end={}
 for no,idx in positions.items():
  # Place tables outside sentences; Round 1 belongs to Appendix A.1.
  anchor=orig[306 if no=='3.2' else 556 if no=='5.2' else 588 if no=='5.5' else 792 if no=='A.2' else idx]
  cap,tab=new_table(doc,anchor,no,payload[no],width);table_end[no]=tab
  remove(orig[idx])
 # Restore fragmented original sentences now that the image interruption is removed.
 for a,b in [(303,305),(559,561),(589,591)]:
  orig[a].add_run(' ')
  for run in list(orig[b]._p):
   if run.tag!=qn('w:pPr'):orig[a]._p.append(copy.deepcopy(run))
  remove(orig[b])
 # Data are restored in Tables A.1/A.3; these are duplicate fragment paragraphs.
 for i in [790,791,*range(796,803)]:remove(orig[i])
 # Add figures using stable original anchors. For removed anchors use the native table.
 for no,idx,path,lead,caption in FIGURES:
  anchor=orig[idx]
  if no=='3.4':
   node=OxmlElement('w:p');table_end['3.1']._tbl.addnext(node);anchor=Paragraph(node,orig[294]._parent)
  elif no=='A.1':
   node=OxmlElement('w:p');table_end['A.3']._tbl.addnext(node);anchor=Paragraph(node,orig[795]._parent)
  pp=para_after(anchor,lead);pp.paragraph_format.keep_with_next=True
  im=para_after(pp);im.alignment=WD_ALIGN_PARAGRAPH.CENTER;im.paragraph_format.keep_with_next=True
  im.add_run().add_picture(str(ROOT/'figures'/path),width=Inches(width))
  cap=para_after(im,f'Figure {no}: {caption}','PDFCaption');cap.paragraph_format.keep_with_next=False;tc(cap,f'{no} {caption.split(".")[0]}.','F')
 # Existing figures retain numbering; add TOC entries without touching research captions.
 for idx,no,title in [(263,'3.1','Hierarchical diagnostic evaluation framework'),(280,'3.2','Task structure and supervision annotation'),(552,'5.1','Validation target-centre frustum coverage'),(565,'5.2','Arm skill relative to no motion'),(597,'5.3','Validation observation-conditioning ratios')]:tc(orig[idx],f'{no} {title}.','F')
 # Explicit in-text references for the pre-existing results figures.
 orig[549].add_run(' Figure 5.1 summarises these geometric coverage ratios by task segment.')
 orig[563].add_run(' Figure 5.2 presents the corresponding arm-skill estimates and inference-seed variation.')
 orig[595].add_run(' Figure 5.3 visualises the perturbation/noise ratios and their anchor-level intervals.')
 orig[19].text='Figures 3.1, 3.2 and 3.4 are schematic diagrams drawn using AI-assisted code. Figure 3.3 contains extracted simulation-video frames, and Figure 3.5 summarises recorded dataset statistics. The statistical'
 # Replace static front-matter indexes with TC-based Word fields. Cached entries omit stale page numbers.
 chapter_pairs=[(77,78),(147,148),(258,259),(451,452),(539,540),(626,627),(717,718),(785,786)]
 entries=[]
 for a,b in chapter_pairs:
  title=orig[a].text.strip()+' '+orig[b].text.strip();tc(orig[a],title,'D',1);entries.append(title)
 for i,p in enumerate(orig):
  if i>=77 and p.style.name in ('SectionTitle','SubsectionTitle') and p._p.getparent() is not None:
   if re.match(r'^(?:[1-7]|A)\.',p.text):tc(p,p.text.strip(),'D',2 if p.style.name=='SectionTitle' else 3)
 orig[35].clear();field(orig[35],'TOC \\f D \\h \\z','\n'.join(entries))
 for i in range(36,49):remove(orig[i])
 figures=[f'{n} {c.split(".")[0]}.' for n,_,_,_,c in FIGURES]
 all_figs=['3.1 Hierarchical diagnostic evaluation framework.','3.2 Task structure and supervision annotation.',*figures[:3],'5.1 Validation target-centre frustum coverage.','5.2 Arm skill relative to no motion.','5.3 Validation observation-conditioning ratios.',figures[-1]]
 orig[70].clear();field(orig[70],'TOC \\f F \\h \\z','\n'.join(all_figs));remove(orig[71])
 orig[73].clear();field(orig[73],'TOC \\f T \\h \\z','\n'.join(f'{no} {p["caption"]}' for no,p in payload.items()))
 for i in (74,75,76):remove(orig[i])
 upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');doc.settings.element.append(upd)
 out.parent.mkdir(parents=True,exist_ok=True);doc.save(out)
 assert hashlib.sha256(inp.read_bytes()).hexdigest()==before
 check=Document(out);assert len(check.tables)==18
 result={'input_sha256':before,'output_sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'tables':len(check.tables),'figures_added':4,'core_figures_total':8,'numbered_figures_including_appendix':9,'original_preserved':True,'paper_size':'A4 (original v6 was US Letter)','page_render_verified':False,'required_final_check':'Open in Microsoft Word, update all fields, then inspect pagination and export PDF.'}
 (ROOT/'reports/DOCX_VALIDATION.json').write_text(json.dumps(result,indent=2));print(json.dumps(result,indent=2))

if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--input',type=Path,required=True);p.add_argument('--output',type=Path,required=True);a=p.parse_args();build(a.input,a.output)
