"""Restore documented equations to native OMML; never writes a document itself.

Call changes = repair_math(doc) before the parent saves/renders its v8 copy.
Every expected original equation fragment is preflighted. A mismatch returns
Pending entries without mutating the document. Native equations are idempotent.
"""
from copy import deepcopy
import re
import unicodedata
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

# Exact v7 text, including source fractions split into consecutive paragraphs.
SOURCE = {
 '3.1': ['s𝑡= bprev 𝑡 , qtorso 𝑡 , q𝐿 𝑡, 𝑔𝐿 𝑡, q𝑅 𝑡, 𝑔𝑅 𝑡 . (3.1)'],
 '3.2': ['Δ𝑎( 𝑗) 𝑡+𝑘= 𝑎( 𝑗) 𝑡+𝑘−𝑠( 𝑗) 𝑡 . (3.2)'],
 '3.3': ['(𝑜𝑡, s𝑡, 𝑝𝑡) →(a𝑡+1, . . . , a𝑡+𝐻) , (3.3)'],
 '3.4': ['𝑃data(𝑑= 1) ≈0.0396, (3.4)'],
 '3.5': ['𝑃train(𝑑= 1) = 0.20, (3.5)'],
 '3.6': ['𝑆= 1 −𝐸𝑚', '𝐸0 . (3.6)'],
 '3.7': ['𝑅motion = E[|Δˆaarm|]', 'E[|Δaarm|] . (3.7)'],
 '3.8': ['ˆ𝑎( 𝑗) 𝑥 −𝑠( 𝑗) 𝑥, 𝑗∈D,', '˜𝑎( 𝑗) 𝑥 =', 'ˆ𝑎( 𝑗) 𝑥, 𝑗∉D, (3.8)'],
 '3.9': ['𝐷(𝑔) 𝑖,𝑥= mean ˜a(𝑔) 𝑥 −˜a(𝑔) ref', '. (3.9)'],
 '3.10': ['𝑄(𝑔) 𝑖,𝑥= 𝐷(𝑔) 𝑖,𝑥', '𝐷(𝑔) 𝑖,noise . (3.10)'],
 '4.1': ['𝑧(𝑥) = 𝑥−𝜇', '𝜎+ 𝜖, (4.1)'],
 '4.2': ['whereas C uses 𝑞(𝑥) = 2 𝑥−𝑞0.01 𝑞0.99 −𝑞0.01 + 𝜖−1. (4.2)'],
}


def _element(tag, *children, val=None):
    node = OxmlElement('m:' + tag)
    if val is not None:
        node.set(qn('m:val'), str(val))
    for child in children:
        node.append(child)
    return node


def _run(text, plain=False):
    node = _element('r')
    if plain:
        node.append(_element('rPr', _element('nor')))
    props = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), 'Cambria Math')
    fonts.set(qn('w:hAnsi'), 'Cambria Math')
    props.append(fonts)
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '24'); props.append(size)
    node.append(props)
    token = _element('t'); token.text = text
    token.set(qn('xml:space'), 'preserve'); node.append(token)
    return node


def _group(tag, nodes):
    return _element(tag, *nodes)


def _var(name, sub=None, sup=None, accent=None, sub_plain=False, sup_plain=False):
    base = _run(name)
    if accent:
        base = _element('acc', _element('accPr', _element('chr', val=accent)), _element('e', base))
    if sub is not None and sup is not None:
        return _element('sSubSup', _element('e', base), _element('sub', _run(sub, sub_plain)), _element('sup', _run(sup, sup_plain)))
    if sub is not None:
        return _element('sSub', _element('e', base), _element('sub', _run(sub, sub_plain)))
    if sup is not None:
        return _element('sSup', _element('e', base), _element('sup', _run(sup, sup_plain)))
    return base


def _frac(numerator, denominator):
    return _element('f', _group('num', numerator), _group('den', denominator))


def _delimit(begin, end, nodes):
    return _element('d', _element('dPr', _element('begChr', val=begin), _element('endChr', val=end)), _group('e', nodes))


def _formula(no):
    v, r = _var, _run
    if no == '3.1':
        components = [('b','prev'),('q','torso'),('q','L'),('g','L'),('q','R'),('g','R')]
        terms = []
        for i, (name, sup) in enumerate(components):
            if i: terms.append(r(', '))
            terms.append(v(name, 't', sup, sup_plain=sup in ('prev','torso')))
        return [v('s','t'),r(' = '),_delimit('(',')',terms)]
    if no == '3.2':
        return [r('Δ'),v('a','t+k','(j)'),r(' = '),v('a','t+k','(j)'),r(' − '),v('s','t','(j)')]
    if no == '3.3':
        return [_delimit('(',')',[v('o','t'),r(', '),v('s','t'),r(', '),v('p','t')]),r(' → '),
                _delimit('(',')',[v('a','t+1'),r(', …, '),v('a','t+H')])]
    if no in ('3.4','3.5'):
        which, sign, value = ('data',' ≈ ','0.0396') if no == '3.4' else ('train',' = ','0.20')
        return [v('P',which,sub_plain=True),_delimit('(',')',[r('d = 1')]),r(sign+value)]
    if no == '3.6':
        return [r('S = 1 − '),_frac([v('E','m')],[v('E','0')])]
    if no == '3.7':
        def magnitude(hat):
            return [r('E'),_delimit('[',']',[_delimit('|','|',[r('Δ'),v('a',sup='arm',accent='̂' if hat else None,sup_plain=True)])])]
        return [v('R','motion',sub_plain=True),r(' = '),_frac(magnitude(True),magnitude(False))]
    if no == '3.8':
        matrix = _element('m')
        for terms, condition in [([v('a','x','(j)',accent='̂'),r(' − '),v('s','x','(j)')], 'j ∈ D'),
                                 ([v('a','x','(j)',accent='̂')], 'j ∉ D')]:
            matrix.append(_element('mr',_group('e',terms),_element('e',r(condition))))
        return [v('a','x','(j)',accent='̃'),r(' = '),_delimit('{','',[matrix])]
    if no == '3.9':
        difference = [v('a','x','(g)',accent='̃'),r(' − '),v('a','ref','(g)',accent='̃',sub_plain=True)]
        return [v('D','i,x','(g)'),r(' = '),r('mean',True),r(' '),_delimit('|','|',difference)]
    if no == '3.10':
        return [v('Q','i,x','(g)'),r(' = '),_frac([v('D','i,x','(g)')],[v('D','i,noise','(g)')])]
    if no == '4.1':
        return [r('z(x) = '),_frac([r('x − μ')],[r('σ + ε')])]
    if no == '4.2':
        return [r('q(x) = 2 '),_frac([r('x − '),v('q','0.01')],[v('q','0.99'),r(' − '),v('q','0.01'),r(' + ε')]),r(' − 1')]
    raise ValueError(no)


def _normal(text):
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', text))


def _superscript_known_powers(doc):
    """Only negative powers already written as 10-minus-n; no numeric changes."""
    changes = []
    pattern = re.compile(r'10([−–-](?:10|[5678]))(?!\d)')
    for paragraph in doc.paragraphs:
        for run in list(paragraph.runs):
            matches = list(pattern.finditer(run.text))
            if not matches or run.font.superscript:
                continue
            if any(child.tag not in (qn('w:rPr'), qn('w:t')) for child in run._r):
                continue  # Preserve fields, hyperlinks, tabs and other run structures.
            text, position, replacement = run.text, 0, []
            def segment(value, superscript=False):
                if not value: return
                node = OxmlElement('w:r')
                props = deepcopy(run._r.rPr) if run._r.rPr is not None else OxmlElement('w:rPr')
                for old in list(props.findall(qn('w:vertAlign'))): props.remove(old)
                if superscript:
                    align = OxmlElement('w:vertAlign');align.set(qn('w:val'),'superscript');props.append(align)
                node.append(props);t=OxmlElement('w:t');t.set(qn('xml:space'),'preserve');t.text=value;node.append(t);replacement.append(node)
            for match in matches:
                segment(text[position:match.start()+2]);segment(match.group(1),True);position=match.end()
            segment(text[position:]);parent=run._r.getparent();at=parent.index(run._r)
            for offset,node in enumerate(replacement):parent.insert(at+offset,node)
            parent.remove(run._r)
            changes.append({'status':'Repaired','kind':'native superscript','original':text,'powers':[m.group(0) for m in matches]})
    return changes


def repair_math(doc):
    """Return auditable changes; edit only the supplied in-memory Document.

    Preflight all 12 original fragment sequences before changing any of them.
    A mismatch gives Pending and leaves doc unchanged. Caller owns saving and
    actual renderer QA; OMML structure alone does not guarantee pagination.
    """
    paragraphs = list(doc.paragraphs);plans=[];changes=[];pending=[]
    for no, originals in SOURCE.items():
        candidates = [(i,p) for i,p in enumerate(paragraphs) if re.search(r'\('+re.escape(no)+r'\)\s*$',p.text)]
        if len(candidates)!=1:
            pending.append({'status':'Pending','equation':no,'reason':f'Expected one numbered equation, found {len(candidates)}'});continue
        index, anchor=candidates[0]
        if anchor._p.xpath('.//m:oMath'):
            changes.append({'status':'Unchanged','equation':no,'reason':'Native OMML already present'});continue
        start=index-len(originals)+1;fragments=paragraphs[max(0,start):index+1]
        if start<0 or [_normal(p.text) for p in fragments] != [_normal(t) for t in originals]:
            pending.append({'status':'Pending','equation':no,'reason':'Original formula fragments do not match; no reconstruction attempted','observed':[p.text for p in fragments]});continue
        plans.append((no,fragments))
    if pending:
        return pending
    for no, fragments in plans:
        paragraph=fragments[0];original=[p.text for p in fragments]
        paragraph.clear();paragraph.style='Normal';paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
        formatting=paragraph.paragraph_format
        formatting.line_spacing=1.2;formatting.space_before=Pt(6);formatting.space_after=Pt(6)
        formatting.keep_together=True;formatting.keep_with_next=False
        if no=='4.2':paragraph.add_run('whereas C uses ').font.size=Pt(12)
        math=_element('oMath',*_formula(no));paragraph._p.append(math)
        paragraph.add_run('    ('+no+')').font.size=Pt(12)
        for extra in fragments[1:]:extra._p.getparent().remove(extra._p)
        changes.append({'status':'Repaired','equation':no,'original':original,'layout':'Native Word OMML; original number retained'})
    return changes+_superscript_known_powers(doc)
