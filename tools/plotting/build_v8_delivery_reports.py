#!/usr/bin/env python3
"""Build v8 revision receipts from the actual final artifacts and QA outputs."""
from pathlib import Path
from zipfile import ZipFile
import copy,csv,difflib,hashlib,json,re,shutil,xml.etree.ElementTree as E
from docx import Document
ROOT=Path(__file__).resolve().parents[2];DEL=Path('/root/dissertation_v8_delivery');AUD=Path('/root/dissertation_local_audit/v8')
V7=Path('/root/dissertation_figures_delivery/NTU_Dissertation_Reviewed_v7_figures_editable.docx');V8=DEL/'NTU_Dissertation_Reviewed_v8_figures_editable.docx';PDF=DEL/'NTU_Dissertation_Reviewed_v8_review.pdf'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
 a=Document(V7);b=Document(V8);q=json.loads((AUD/'final_pdf_qa/pdf_qa.json').read_text());s=json.loads((ROOT/'reports/V8_DOCX_STRUCTURE.json').read_text())
 assert not q['flags'] and s['passed'];assert sha(V7)=='e24f23f635b43b97297273ae1f85b66d97b2b09a39e971dee9725b76694afbb9'
 vals=lambda d:[[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
 assert vals(a)==vals(b) and len(b.tables)==18 and len(b.element.xpath('.//m:oMath'))==12
 refs=lambda d:[p.text for p in d.paragraphs if re.match(r'^\[\d+\]',p.text)]
 assert refs(a)==refs(b)
 meta={'version':'v8','date':'2026-10-22 Asia/Singapore','author':'WUHAODONG','v7_sha256':sha(V7),'v7_preserved':True,'docx':{'file':V8.name,'bytes':V8.stat().st_size,'sha256':sha(V8)},'pdf':{'file':PDF.name,'bytes':PDF.stat().st_size,'sha256':sha(PDF),'pages':q['pages'],'renderer':'LibreOffice 7.3.7.2 Writer + Math'},'native_tables':18,'table_values_preserved':True,'native_omml_equations':12,'reference_text_preserved':True,'figure_numbers':s['figure_numbers'],'event_number':'5.3','conditioning_number':'5.4','official_template_status':'Pending official template verification','redistribution':'Full manuscript and recorded-frame storyboard retained locally because underlying scene/human rendering redistribution permission is not established.'}
 (ROOT/'manuscript').mkdir(exist_ok=True);(ROOT/'manuscript/V8_MANUSCRIPT_METADATA.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2));(DEL/'reports').mkdir(exist_ok=True);(DEL/'tables').mkdir(exist_ok=True)
 shutil.copy2(AUD/'final_pdf_qa/pdf_qa.json',ROOT/'reports/V8_PDF_QA.json')
 info={
 '3.1':('Hierarchy','methodology/v8_fig_3_1_diagnostic_hierarchy','Revised','V8_METHOD_FIGURES.md'),
 '3.2':('Task and shift-1 supervision','methodology/v8_fig_3_2_task_supervision','Revised','V8_METHOD_FIGURES.md'),
 '3.3':('Six actual simulation frames','methodology/v8_fig_3_3_task_storyboard','Revised layout; local only','tables/storyboard_frames.json'),
 '3.4':('Scripted scheduler and controller boundary','methodology/v8_fig_3_4_controller_data_flow','Revised','V8_METHOD_FIGURES.md'),
 '3.5':('QC and sampling target','dataset/fig_3_5_dataset_construction','Retained; layout repaired','tables/dataset_checks.json'),
 '3.6':('Temporal/action characteristics','dataset/fig_3_6_temporal_action_characteristics','New','V8_EVIDENCE_FIGURES.md'),
 '5.1':('Validation frustum coverage','results/v8_fig_5_1_frustum','Revised; source axis clipping repaired','V8_FRUSTUM_FIGURE.md'),
 '5.2':('Arm skill relative to no motion','results/v8_fig_5_2_arm_skill','Retained','evidence/frozen_audit/comparison.csv'),
 '5.3':('Gripper event profile','results/fig_5_4_gripper_event_performance','New; requested as 5.4','V8_EVIDENCE_FIGURES.md'),
 '5.4':('Observation conditioning','results/v8_fig_5_4_conditioning','Retained; renamed from 5.3','evidence/conditioning/summary.json'),
 'A.1':('Training diagnostics','experiments/fig_A_1_training_diagnostics','Retained; layout repaired','tables/training_curves.csv')}
 mapping=[]
 for no in s['figure_numbers']:
  title,stem,status,source=info[no];p=q['figures'][no]['page'];row={'number':no,'title':title,'status':'Generated','change':status,'pdf_page':p,'file':'figures/'+stem+'.png','source':source,'local_only':no=='3.3','caption_rendered_same_page':True,'requested_event_number':'5.4' if no=='5.3' else None};mapping.append(row)
  for ext in ['png','pdf','svg']:
   src=ROOT/'figures'/f'{stem}.{ext}'
   if src.exists():
    dest=DEL/'figures'/f'{stem}.{ext}';dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(src,dest)
 # Extract the two unchanged result rasters from the actual final Word relationships.
 n={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
 with ZipFile(V8) as z:
  dr=E.fromstring(z.read('word/document.xml'));rr=E.fromstring(z.read('word/_rels/document.xml.rels'));rel={e.get('Id'):e.get('Target') for e in rr};ps=dr.findall('w:body/w:p',n)
  for idx,p in enumerate(ps):
   tx=''.join(t.text or '' for t in p.findall('.//w:t',n));style=p.find('w:pPr/w:pStyle',n)
   if style is None or 'index' in style.get('{'+n['w']+'}val','').lower():continue
   mt=re.match(r'^Figure (5\.[24]):',tx)
   if not mt:continue
   no=mt.group(1)
   for old in reversed(ps[:idx]):
    im=old.find('.//a:blip',n)
    if im is not None:
     import posixpath
     path=posixpath.normpath('word/'+rel[im.get('{'+n['r']+'}embed')]);dest=DEL/'figures'/f'{info[no][1]}.png';dest.parent.mkdir(parents=True,exist_ok=True);dest.write_bytes(z.read(path));break
 (ROOT/'tables/v8_figure_to_thesis_mapping.json').write_text(json.dumps(mapping,ensure_ascii=False,indent=2));(ROOT/'tables/figure_to_thesis_mapping.json').write_text(json.dumps(mapping,ensure_ascii=False,indent=2))
 with (ROOT/'tables/figure_to_thesis_mapping.csv').open('w',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(mapping[0]));w.writeheader();w.writerows(mapping)
 figure_table='\n'.join(f'|{x["number"]}|{x["title"]}|{x["change"]}|{x["pdf_page"]}|`{x["file"]}`|`{x["source"]}`|' for x in mapping)
 common='''日期：2026-10-22（Asia/Singapore）。本轮只修订v7，没有重新训练、仿真、数据集生成或重复完整研究审计。事件图原请求为Figure5.4；为保持首次出现顺序，正文用Figure5.3，原conditioning图顺延5.4。对应脚本stem保留原请求名，映射明确。
'''
 validation=f'''# V8 document validation\n\n{common}\n| Check | Status | Actual evidence |\n|---|---|---|\n| v7 preserved | Pass | SHA256 `{sha(V7)}`; local backup same hash |\n| DOCX ZIP/XML and image relationships | Pass | V8_DOCX_STRUCTURE.json; no missing drawing links |\n| Numbered figures / native tables | Pass | 11 figures, 18 native tables; exact chapter sequences checked in actual PDF |\n| Table values / references | Pass | All 18 tables and numbered-reference text equal v7 |\n| Images and captions | Pass | All11 complete and same-page in LibreOffice render; no srcRect crop/exact picture line-height |\n| PDF all-page checks | Pass | {q['pages']} pages scanned, no text/image outside page, blank page, or caption without same-page image |\n| Native mathematics | Pass | 12 OMML equations, fractions/cases restored from original definitions; actual rendered content checked |\n| Page layout / indexes | Pass | A4,3 sections; cover no numbers, front Roman, body Arabic; 11/18 complete index entries |\n| Official EEE/July2026 declaration | Pending | Pending official template verification; preserve real disclosure and signature placeholders |\n| Microsoft Word rendering equivalence | Pending | PDF actually rendered with LibreOffice; final author Word preview still recommended |\n\nDOCX SHA256 `{sha(V8)}`; PDF SHA256 `{sha(PDF)}`. This is a real review PDF, not an institutional submission certification. Font fallback: Times New Roman maps to Liberation Serif and mathematical text to STIXGeneral on this Linux renderer; Word fonts may lay out differently.\n'''
 (ROOT/'reports/V8_DOCUMENT_VALIDATION.md').write_text(validation)
 visual=f'''# V8 visual QA\n\n{common}\n**Final PDF: {q['pages']} pages, SHA256 `{sha(PDF)}`.**\n\nPass: every page was rasterized and scanned for geometry, near-edge text, image bounds and blank pages; every page was visually reviewed on contact sheets. Figure pages and table/math problem pages were opened at individual-page scale; separate reviewer reports record inspected ranges and intermediate candidate hashes. Only the final hash above governs this report.\n\nResolved: fixed picture-line-height clipping (v7 p26 narrow-strip storyboard versus auto-spacing probe p27 full storyboard); congested method diagrams; incomplete private-TC indexes; cross-chapter SEQ accumulation (explicit per-caption resets and exact PDF label checks); run-on acronyms; short table5.5 splitting; TableA8 Median header wrapping; source-image5.1 y-axis clipping; inherited flattened math; formula renderer's initially missing Math component. The latter was caught before publication and verified with a12-equation probe plus actual final pages.\n\n| Figure | Purpose | Change | PDF page | File | Source |\n|---|---|---|---|\n{figure_table}\n\nTableA3 continues across pages with repeated header and all31 data rows; no fixed row clipping. Short table5.5 is kept together. No unresolved Fail remains in the checked output. A small section opening near the bottom of p32 has two following lines and is not a detached heading; a stylistic author review may still prefer a different break.\n\nPending: Microsoft Word cross-engine preview and exact NTU declaration version/applicability; no generated signature or claim of official approval. Source frame redistribution remains unconfirmed, so full manuscript/contact images stay local.\n'''
 (ROOT/'reports/V8_VISUAL_QA.md').write_text(visual)
 trace=f'''# V8 figure traceability\n\n{common}\n| Figure | Purpose | Change | PDF page | File | Source |\n|---|---|---|---|\n{figure_table}\n\nMethod diagrams: `generate_v8_methodology.py` and V8_METHOD_FIGURES.md. Storyboard: same six video-frame files/identifiers from tables/storyboard_frames.json, relaid at5.7in by generate_v8_storyboard.py; no invented pixels or policy-success claim. Frustum5.1: generate_v8_frustum.py plus v8_frustum_coverage.csv and V8_FRUSTUM_FIGURE.md. Dataset3.6 and event profile: generate_v8_evidence.py plus v8_* CSV/JSON and V8_EVIDENCE_FIGURES.md. All new/revised graphics havePNG/PDF/SVG and editable plotting sources. The two unchanged result figures retain their raster originals without artificial vector conversion.\n\n3.6 populations:45 sources,495 actual non-standby segments; per-source/context mean next-action deviation over18 verified revolute joints (rad); raw frames versus actual B/C valid masks (409670 versus387544). Segments within source are correlated; all splits are descriptive. Raw/anchor units remain separate.\n\nEvent profile:15 per-seed frozen summaries,652 total anchors; pickup/radio each200,4 validation sources. Mean±sample SD over five inference noise seeds; one20k training run per configuration. F1 and precision agree with native Table5.3; no grasp-success, single-factor causal, or significance claim.\n\nCommands: `python tools/plotting/generate_v8_methodology.py`; `python tools/plotting/generate_v8_evidence.py --from-tables`; local storyboard requires `--frames /path/to/verified/frames`. Neither full dataset nor weights are needed for from-tables plotting.\n'''
 (ROOT/'reports/V8_FIGURE_TRACEABILITY.md').write_text(trace)
 changes=f'''# V8 academic changelog\n\n{common}\n- Redrew3.1/3.2/3.4, retained six actual frames for3.3 with larger labels, retained3.5 QC/20% target meaning.\n- Added3.6 after§3.3's RQ1 discussion and event profile after§5.3's interpretation, supporting RQ1/RQ2 without changing study conclusions.\n- Adjusted event/conditioning numbering to5.3/5.4 and all captions, references and indexes.\n- Corrected Figure5.1 reference to validation-only segment coverage rather than an all-source antecedent, and redrew from the unchanged archived counts to repair source-image axis-label clipping.\n- Restored12 equations as native OMML using original fragments; no parameter or formula meaning changed.\n- Fixed display geometry, native heading/SEQ indexes, Roman/Arabic footers, acronyms and table pagination; all18 table values and references preserved.\n- Updated truthful AI assistance/figure-source disclosure, preserving declarations and unsigned placeholders. Official EEE template remains Pending.\n- Retained45 sources/37-4-4 split/652 validation anchors/five inference seeds/one training run/108validation+216train-coverage distinction/historical rollout boundary.\n- No new policy inference, training, rollout success, dataset-scale or generalization claim.\n\nFull paragraph/XML-text diff and v7 backup are local delivery artifacts; the repository receives this concise change list and manuscript hashes, not embedded restricted scene imagery.\n'''
 (ROOT/'reports/V8_ACADEMIC_CHANGELOG.md').write_text(changes)
 # Private, author-review-only text difference; not added to repository.
 text=lambda d:[p.text for p in d.paragraphs]
 (DEL/'V7_TO_V8_DIFF.txt').write_text(''.join(difflib.unified_diff([x+'\n' for x in text(a)],[x+'\n' for x in text(b)],fromfile='v7',tofile='v8')))
 receipt_path=ROOT/'reports/V8_PUSH_VERIFICATION.json'
 if receipt_path.exists():
  receipt=json.loads(receipt_path.read_text());assert receipt['passed'] and receipt['private']
  sync_info=f"**Pass**：branch `{receipt['branch']}`；内容提交 `{receipt['remote_commit_sha']}`。GitHub API确认Private，fresh fetch后的远程tree与本地一致，全部{receipt['tracked_files_verified']}个文件逐一SHA256一致；无受禁资产。该回执验证内容提交；最终补记报告的HEAD另见本地V8_RELEASE_RECEIPT.json。更新文件的完整列表及hash见V8_PUSH_VERIFICATION.json。"
 else:
  sync_info='内容已准备，push尚待执行；完成后本生成器会读取V8_PUSH_VERIFICATION.json写入实际核验。'
 final=f'''# v8最终交付报告\n\n{common}\n## A. What Changed\n修复四张图继承exact行高导致的裁切；重绘3.1/3.2/3.4及5.1（修复旧图纵轴标签裁切），重排同六帧3.3标签；新增3.6和事件profile（正文5.3，原请求5.4），原conditioning改5.4。保留原章节、题目、WUHAODONG、18表数据、References、Appendix和签字占位。恢复12式为原生Word数学，补全可更新目录与页码，修复缩写和表格分页。\n\n## B. Figures\n| Figure | Purpose | Change | PDF page | File | Source |\n|---|---|---|---|\n{figure_table}\n\n11图均在实际PDF中完整显示并与caption同页。新增/重绘图提供PNG/PDF/SVG和源脚本。\n\n## C. Dataset and Benchmark Improvements\n3.6用真实分段时长、18维旋转关节next-action偏差和raw/valid总体区别解释长时程监督失衡，支撑RQ1；495个segments不是495个独立sources。事件图将pickup/radio的F1、precision及inference variability并置，支撑RQ2的窄结论；没有把事件命中写成抓取成功或模型总体能力。\n\n## D. DOCX/PDF QA\n实际LibreOffice Writer+Math渲染{q['pages']}页，全部页完成几何扫描和contact概览，重点图/数学/表格单页放大核查。零越界/异常空白/图注无同页图；全部11图和18表号严格核对，A3续页表头完整、5.5整表同页、A8表头不再断词。Word有12原生OMML与18原生表，原v7哈希保持。详见V8_DOCUMENT_VALIDATION.md、V8_VISUAL_QA.md。\n\nDOCX：`{V8}`，{V8.stat().st_size:,}bytes，SHA256`{sha(V8)}`。\n\nPDF：`{PDF}`，{PDF.stat().st_size:,}bytes，SHA256`{sha(PDF)}`。\n\n## E. Academic Integrity\n全部数值来自已有记录或本轮真实派生统计；没有生成仿真证据、训练或成功率。保留scripted FSM、upstream replay、Execute内低层policy的贡献边界。五seed是推理噪声，A/B与B/C均非单因素实验；teacher forcing和历史closed-loop分开。学校具体声明措辞未擅造，未代签。\n\n## F. GitHub Sync\n目标私有仓库https://github.com/lucaswhd/dissertation，开始时已fetch且main与9c2cf27一致。上传非视频图、代码、CSV/JSON、README、映射与报告；完整Word/PDF、真实帧和逐页预览仅本地（第三方渲染再分发许可未明确）。{sync_info}\n\n## G. Remaining NTU Submission Actions\n1. **Pending official template verification**：已取得真实25Nov2025公开NTU模板，尚未查实July2026后EEE适用版本；由学校/作者核对，保留真实AI披露并补确知的工具版本。\n2. 作者在Microsoft Word做最终跨引擎预览、更新域；Linux实际PDF已审核，但字体替代可能令Word分页略变。\n3. 作者/导师本人签字和日期；本轮未代填。\n4. 完整稿/帧向仓库或外部再分发前核对上游场景/人物渲染许可；本轮已隔离，未阻止安全材料发布。\n'''
 (ROOT/'reports/V8_FINAL_DELIVERY_REPORT.md').write_text(final)
 (DEL/'V8_DELIVERY_README.md').write_text('''# v8交付使用说明

打开NTU_Dissertation_Reviewed_v8_figures_editable.docx；同目录PDF为已实际渲染并逐页检查的78页审阅版。原v7备份在backup/，文字差异在V7_TO_V8_DIFF.txt。

事件图为Figure5.3，原conditioning顺延5.4（按首次出现顺序）；需求中的event5.4源文件stem保留，映射表说明此差异。18表可编辑，12式为原生Word数学，新增/重绘图有矢量源。

Word可Ctrl+A→F9更新目录/图表目录/页码。现有题注序号显式按章固定，避免不同渲染器跨章累计；未来新增图时应在生成器映射中调整编号并同步引用。提交前请在Word检查字体/分页，再核对学校最新版EEE声明、真实AI工具版本及本人/导师签字。

完整论文和真实视频帧只在本地交付，未上传Github；安全的图表代码、数值统计和报告同步状态见reports/V8_FINAL_DELIVERY_REPORT.md。最终推送记录见reports/V8_PUSH_VERIFICATION.json与V8_RELEASE_RECEIPT.json（完成后生成）。
''')
 # Copy all current v8-only reports and tables; copy again after push for its receipts.
 for p in (ROOT/'reports').glob('V8*'):
  if p.is_file():shutil.copy2(p,DEL/'reports'/p.name)
 for p in (ROOT/'tables').glob('v8_*'):
  if p.is_file():shutil.copy2(p,DEL/'tables'/p.name)
 (DEL/'tables/figure_to_thesis_mapping.json').write_text(json.dumps(mapping,ensure_ascii=False,indent=2))
 print(json.dumps({'docx_sha256':sha(V8),'pdf_sha256':sha(PDF),'pages':q['pages'],'figures':len(mapping),'tables_values_preserved':True,'native_equations':12,'reports_written':5},indent=2))
if __name__=='__main__':main()
